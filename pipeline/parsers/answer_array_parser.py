"""Parser for the wide-format ``Массив ответов`` (answer array) workbooks.

Layout of every sheet (one sheet == one test):

    Row 0 : Пользователь | Дата | Статус | Баллы | <Q1 text> | "" | "" | "" | <Q2 text> | ...
            (each question heading is merged across its 4 columns, so only the
             first column of the group carries text)
    Row 1 :   ""   |  ""  |   ""   |  ""   | Тип | Результат | Полученный ответ | Правильный ответ | Тип | ...
    Row 2+: data rows, one per attempt.

Key robustness decision: within each question's column block we locate the four
sub-columns by their **Row-1 labels** (Тип / Результат / Полученный ответ /
Правильный ответ) rather than assuming a fixed +0..+3 offset. Merged cells and
the occasional ragged block can shift offsets; label matching cannot misalign.

A question block that is entirely empty for a given row means that question was
not shown to that student (question-bank randomisation) and is skipped.

Hardening (all deterministic, zero tokens):

* the sub-header row (Тип / Результат / ...) is searched in the first rows of the
  sheet, so banner / title rows above the table are tolerated;
* the leading columns (Пользователь / Дата / Статус / Баллы) are found by their
  labels, with the classic positions 0..3 as a fallback; extra leading columns are
  ignored;
* a sheet that cannot be read is reported through ``notes`` instead of being
  silently skipped;
* result codes: the LMS hashes, textual verdicts (верно/неверно/...) and - as a last
  resort - a code inferred from ``given == correct`` on choice questions, always
  with a loud warning.
"""

from __future__ import annotations

import logging
import re
import uuid
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from core.exceptions import ParseError
from core.models import (
    QuestionAnswer,
    QuestionType,
    ResultCode,
    StudentAttempt,
)
from pipeline.parsers.excel_utils import (
    coerce_datetime,
    is_service_sheet,
    norm_cell,
    parse_score,
    unreadable_file_message,
)

log = logging.getLogger("korpus.parser")

# Classic leading columns (positional fallback when labels are not found).
_FIXED_COLS = 4
_COL_USER, _COL_DATE, _COL_STATUS, _COL_SCORE = range(_FIXED_COLS)

#: How many top rows are searched for the sub-header row (banner-row tolerance).
_HEADER_SEARCH_ROWS = 25

# Leading-column labels -> role (matched on the casefolded cell, first hit wins).
_LEADING_PATTERNS: tuple[tuple[str, re.Pattern], ...] = (
    # «Результат попытки» / «Итог попытки» is the pass/fail STATUS, not the score.
    ("status", re.compile(r"^(статус|status|состоян|(результат|итог|исход)\s+(попытки|теста|тестирования|прохождения)|зачет|зачёт)")),
    ("score", re.compile(r"^(балл|набрано|оценк|score|итог|процент|результат)")),
    ("user", re.compile(
        r"^(пользоват|фио|ф\.и\.о|логин|студент|сотрудник|учащ|слушател|участник|код|табельн|номер|идентиф|id$|e-?mail|почта|user|login|name|имя)"
    )),
    ("date", re.compile(r"^(дата|date|начал|заверш|прохожд|время(?!.*(мин|сек|длит|duration)))")),
)

# Row-1 sub-header label fragments -> logical role. Matched case-insensitively
# by prefix so "Полученный ответ" and "Правильный ответ" never collide.
_SUBHEADER_MATCHERS: tuple[tuple[str, str], ...] = (
    ("тип", "type"),
    ("результат", "result"),
    ("получ", "given"),
    ("правильн", "correct"),
)


@dataclass(frozen=True)
class _QuestionLayout:
    """Resolved column positions for one question within a sheet."""

    text: str
    type_col: int
    result_col: int
    given_col: int
    correct_col: int


def _subheader_role(label: str) -> str | None:
    for fragment, role in _SUBHEADER_MATCHERS:
        if label.startswith(fragment):
            return role
    return None


def _subheader_score(row: tuple) -> int:
    """Number of distinct question sub-column roles present in ``row``."""
    return len({r for c in row if (r := _subheader_role(norm_cell(c).casefold()))})


def _find_header_rows(rows: list[tuple]) -> int | None:
    """Index of the sub-header row (Тип/Результат/...), searched in the top rows.

    The row directly above it carries the question headings. ``None`` when no row
    in the search window has at least three of the four sub-column roles."""
    for i, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        if _subheader_score(row) >= 3:
            return i
    return None


def _first_block_col(header1: tuple) -> int:
    for col, c in enumerate(header1):
        if _subheader_role(norm_cell(c).casefold()):
            return col
    return _FIXED_COLS


def _resolve_leading(header0: tuple, header1: tuple, first_block: int) -> dict[str, int]:
    """Find user/date/status/score columns by label before the first question block.

    Labels may sit in either header row. Missing roles fall back to the classic
    positions 0..3 when those still lie in the leading area and are unclaimed."""
    found: dict[str, int] = {}
    width = max(len(header0), len(header1))
    for col in range(min(first_block, width)):
        for hdr in (header0, header1):
            label = norm_cell(hdr[col]).casefold() if col < len(hdr) else ""
            if not label:
                continue
            hit = next((role for role, pat in _LEADING_PATTERNS if pat.match(label)), None)
            if hit is not None and hit not in found:
                found[hit] = col
                break
    claimed = set(found.values())
    for role, pos in (("user", _COL_USER), ("date", _COL_DATE),
                      ("status", _COL_STATUS), ("score", _COL_SCORE)):
        if role not in found and pos < first_block and pos not in claimed:
            found[role] = pos
            claimed.add(pos)
    if "user" not in found:  # last resort: the first leading column nobody claimed
        free = [c for c in range(min(first_block, width)) if c not in claimed]
        if free:
            found["user"] = free[0]
    return found


def _resolve_question_layouts(
    header0: tuple, header1: tuple, *, filepath: str, sheet: str, first_block: int = _FIXED_COLS
) -> list[_QuestionLayout]:
    """Locate each question block and its four sub-columns via Row-1 labels."""
    width = max(len(header0), len(header1))
    header0 = tuple(header0) + (None,) * (width - len(header0))
    # Starts: columns at/after the leading block whose heading cell is non-empty.
    starts = [c for c in range(first_block, width) if norm_cell(header0[c])]
    if not starts:
        return []

    layouts: list[_QuestionLayout] = []
    boundaries = starts + [width]
    for idx, start in enumerate(starts):
        end = boundaries[idx + 1]
        roles: dict[str, int] = {}
        for col in range(start, end):
            label = norm_cell(header1[col]).casefold() if col < len(header1) else ""
            if not label:
                continue
            role = _subheader_role(label)
            if role and role not in roles:
                roles[role] = col
        # Fall back to fixed offsets for any role the labels didn't pin down.
        layouts.append(
            _QuestionLayout(
                text=norm_cell(header0[start]),
                type_col=roles.get("type", start),
                result_col=roles.get("result", start + 1),
                given_col=roles.get("given", start + 2),
                correct_col=roles.get("correct", start + 3),
            )
        )
    return layouts


def _cell(row: tuple, col: int) -> object:
    return row[col] if 0 <= col < len(row) else None


def _build_answers(
    row: tuple, layouts: list[_QuestionLayout], raw_codes: dict[int, str] | None = None
) -> list[QuestionAnswer]:
    answers: list[QuestionAnswer] = []
    for q in layouts:
        type_raw = norm_cell(_cell(row, q.type_col))
        result_raw = norm_cell(_cell(row, q.result_col))
        given = norm_cell(_cell(row, q.given_col))
        correct = norm_cell(_cell(row, q.correct_col))
        # Entirely empty block => question not shown to this student.
        if not (type_raw or result_raw or given or correct):
            continue
        qa = QuestionAnswer(
            question_text=q.text,
            question_type=QuestionType.from_raw(type_raw),
            result_code=ResultCode.from_raw(result_raw),
            given_answer=given,
            correct_answer=correct,
            is_answered=bool(given),
        )
        if qa.result_code is None and result_raw and raw_codes is not None:
            raw_codes[id(qa)] = result_raw
        answers.append(qa)
    return answers


def _note(notes: list[str] | None, msg: str) -> None:
    log.warning(msg)
    if notes is not None:
        notes.append(msg)


def _infer_result_codes(
    attempts: list[StudentAttempt], raw_codes: dict[int, str], notes: list[str] | None
) -> None:
    """Unknown result tokens: learn their meaning from ``given == correct`` on choice
    questions. Applied only when a token is overwhelmingly consistent; always warns."""
    answers_by_id = {id(a): a for at in attempts for a in at.answers}
    by_token: dict[str, list[QuestionAnswer]] = defaultdict(list)
    for aid, tok in raw_codes.items():
        if aid in answers_by_id:
            by_token[tok].append(answers_by_id[aid])
    for tok, answers in by_token.items():
        sample = [
            a for a in answers
            if a.question_type in (QuestionType.SINGLE_CHOICE, QuestionType.MULTI_CHOICE)
            and a.given_answer and a.correct_answer
        ]
        verdict = None
        if len(sample) >= 3:
            same = sum(
                1 for a in sample
                if a.given_answer.strip().casefold() == a.correct_answer.strip().casefold()
            )
            frac = same / len(sample)
            verdict = (
                ResultCode.CORRECT if frac >= 0.9 else ResultCode.WRONG if frac <= 0.1 else None
            )
        if verdict is not None:
            for a in answers:
                a.result_code = verdict
            word = "ВЕРНО" if verdict is ResultCode.CORRECT else "НЕВЕРНО"
            _note(notes, f"Неизвестный код результата «{tok}» ({len(answers)} ответов): определён "
                         f"как {word} по совпадению полученного и правильного ответа. "
                         f"Проверьте колонку «Результат».")
        else:
            _note(notes, f"ВНИМАНИЕ: неизвестный код результата «{tok}» ({len(answers)} ответов) - "
                         f"эти ответы НЕ учитываются как верные/неверные. "
                         f"Ожидаются коды LMS либо «верно/неверно».")


def parse_sheet_rows(
    rows: list[tuple],
    sheet_name: str,
    *,
    course_id: str,
    filepath: str = "",
    notes: list[str] | None = None,
    raw_codes: dict[int, str] | None = None,
) -> list[StudentAttempt]:
    """Parse one sheet-like grid (xlsx sheet or a wide CSV) into attempts.

    Shared by the xlsx and the wide-CSV paths so both get the same label-based column
    resolution, banner tolerance and result-code handling."""
    attempts: list[StudentAttempt] = []
    if raw_codes is None:
        raw_codes = {}
    if not any(norm_cell(c) for r in rows for c in r):
        return []  # genuinely empty sheet: nothing to report
    if is_service_sheet(sheet_name):
        return []  # «Сводка» / «Информация» / ...: not a test
    hi = _find_header_rows(rows)
    if hi is None or hi < 1:
        _note(notes, f"Лист «{sheet_name}» пропущен: не найдена строка заголовков "
                     f"вопросов (Тип / Результат / Полученный ответ / Правильный ответ).")
        return attempts
    header0, header1 = rows[hi - 1], rows[hi]
    first_block = _first_block_col(header1)
    layouts = _resolve_question_layouts(
        header0, header1, filepath=filepath, sheet=sheet_name, first_block=first_block
    )
    if not layouts:
        _note(notes, f"Лист «{sheet_name}» пропущен: не найдено ни одного вопроса.")
        return attempts
    lead = _resolve_leading(header0, header1, first_block)
    if "user" not in lead:
        _note(notes, f"Лист «{sheet_name}» пропущен: не найдена колонка «Пользователь».")
        return attempts
    for role, label in (("date", "Дата"), ("status", "Статус"), ("score", "Баллы")):
        if role not in lead:
            _note(notes, f"Лист «{sheet_name}»: не найдена колонка «{label}» - "
                         f"значения не заполнены.")
    data_rows = rows[hi + 1:]
    if not data_rows:
        _note(notes, f"Лист «{sheet_name}» пропущен: нет строк с данными.")
        return attempts
    c_user = lead["user"]
    c_date = lead.get("date", -1)
    c_status = lead.get("status", -1)
    c_score = lead.get("score", -1)

    for row in data_rows:
        student_id = norm_cell(_cell(row, c_user))
        if not student_id:
            continue  # blank/spacer row
        points, maximum, pct = parse_score(_cell(row, c_score))
        answers = _build_answers(row, layouts, raw_codes)
        answered_any = any(a.is_answered for a in answers)
        is_abandoned = pct == 0.0 and not answered_any
        attempts.append(
            StudentAttempt(
                attempt_id=str(uuid.uuid4()),
                student_id=student_id,
                course_id=course_id,
                test_name=sheet_name,
                attempt_date=coerce_datetime(_cell(row, c_date)),
                score_points=points,
                score_max=maximum,
                score_pct=pct,
                status=norm_cell(_cell(row, c_status)),
                is_abandoned=is_abandoned,
                answers=answers,
            )
        )

    return attempts


def parse_answer_array(
    filepath: str | Path, course_id: str, notes: list[str] | None = None
) -> list[StudentAttempt]:
    """Parse one ``Массив ответов`` workbook into a flat list of attempts.

    Every sheet is treated as a separate test; ``test_name`` is the sheet name.
    Non-fatal findings (skipped sheets, unknown result codes, ...) are appended to
    ``notes`` (and logged) instead of being swallowed.
    """
    path = Path(filepath)
    if not path.exists():
        raise ParseError("Файл не найден", filepath=str(path))

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a variety of types
        raise ParseError(unreadable_file_message(path.name, exc), filepath=str(path)) from exc

    attempts: list[StudentAttempt] = []
    raw_codes: dict[int, str] = {}  # id(QuestionAnswer) -> unknown raw result token
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 3:
                try:  # read-only mode trusts a possibly wrong <dimension>; re-measure
                    ws.reset_dimensions()
                    rows = list(ws.iter_rows(values_only=True))
                except Exception:  # noqa: BLE001
                    pass
            attempts.extend(parse_sheet_rows(
                rows, sheet_name, course_id=course_id, filepath=str(path),
                notes=notes, raw_codes=raw_codes,
            ))
    finally:
        wb.close()

    if raw_codes:
        _infer_result_codes(attempts, raw_codes, notes)
    return attempts
