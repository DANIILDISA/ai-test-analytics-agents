"""Parser for the ``Эталон ответов`` (reference answer key) workbooks.

Layout of every sheet (one sheet == one test):

    Row 0 : <Q1 text> | <Q2 text> | <Q3 text> | ...   (one column per question)
    Row 1 : Правильный ответ | Правильный ответ | ...
    Row 2+: the canonical answer(s) for each question.

A single cell may hold several acceptable variants. Per the spec these are
separated by ``", "`` or ``";\\n"``; however a legitimate answer can itself
contain a comma, so by default we only split on the *unambiguous* separators
(``;`` and newlines) and keep ``raw_answer`` intact. The comparison layer can
split more aggressively when it has the question type in hand.

Questions are keyed by normalised text because the Эталон has no IDs and its
question order does not match the Массив (linking happens later, by text).
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

import openpyxl

from core.exceptions import ParseError
from core.models import ReferenceAnswer
from pipeline.parsers.excel_utils import (
    is_service_sheet,
    norm_cell,
    normalize_question,
    unreadable_file_message,
)

# Unambiguous variant separators: ';' (optionally followed by newline) and bare
# newlines. Comma-splitting is deliberately left to the comparison layer.
_VARIANT_SPLIT_RE = re.compile(r";\s*\n?|\n+")

_LABEL_RE = re.compile(r"^(правильн\w*\s+ответ\w*|эталон\w*)$")


def split_variants(raw: str) -> list[str]:
    """Split a reference cell into acceptable variants on safe separators."""
    parts = [p.strip() for p in _VARIANT_SPLIT_RE.split(raw)]
    return [p for p in parts if p]


_SEARCH_ROWS = 25
_log = logging.getLogger("korpus.parser")


def _row_has_label(row: tuple) -> bool:
    return any(_LABEL_RE.match(norm_cell(c).casefold()) for c in row)


def _row_blank(row: tuple) -> bool:
    return not any(norm_cell(c) for c in row)


_QUESTION_HEAD_RE = re.compile(r"^(вопрос\w*|question\w*|задани\w*|текст вопроса)$")


def find_transposed_layout(rows: list[tuple]) -> tuple[int, int, int] | None:
    """Transposed key: one question per ROW. Returns ``(header_row, question_col,
    answer_col)`` for a sheet whose header row has «Вопрос» and «Правильный ответ»
    columns (any extra columns such as «№» are ignored), else ``None``."""
    for i, row in enumerate(rows[:_SEARCH_ROWS]):
        q_col = a_col = None
        for c, cell in enumerate(row):
            label = " ".join(norm_cell(cell).casefold().split())
            if q_col is None and _QUESTION_HEAD_RE.match(label):
                q_col = c
            elif a_col is None and _LABEL_RE.match(label):
                a_col = c
        if q_col is not None and a_col is not None:
            return i, q_col, a_col
    return None


def _parse_transposed(rows, layout, sheet_name):
    h, q_col, a_col = layout
    sheet_map: dict[str, ReferenceAnswer] = {}
    for r in rows[h + 1:]:
        q = norm_cell(r[q_col]) if q_col < len(r) else ""
        a = norm_cell(r[a_col]) if a_col < len(r) else ""
        if not q or not a:
            continue
        key = normalize_question(q)
        if key in sheet_map:  # same question listed twice: keep every variant
            prev = sheet_map[key]
            sheet_map[key] = ReferenceAnswer(
                test_name=sheet_name, question_text=prev.question_text,
                raw_answer=prev.raw_answer + " | " + a,
                acceptable=prev.acceptable + split_variants(a),
            )
            continue
        sheet_map[key] = ReferenceAnswer(
            test_name=sheet_name, question_text=q, raw_answer=a, acceptable=split_variants(a)
        )
    return sheet_map


def find_reference_layout(rows: list[tuple]) -> tuple[int, int] | None:
    """Return ``(header_row, first_answer_row)`` for a reference sheet, or ``None``.

    The "Правильный ответ" label row is searched in the top rows (banner rows above
    the table are tolerated); the question headings are the nearest non-blank row above
    it. Falls back to the classic layout (headings in row 0, answers from row 1)."""
    for i, row in enumerate(rows[:_SEARCH_ROWS]):
        if i >= 1 and _row_has_label(row):
            h = i - 1
            while h > 0 and _row_blank(rows[h]):
                h -= 1
            return h, i + 1
    if len(rows) >= 2 and not _row_blank(rows[0]):
        return 0, 1
    return None


def parse_reference(
    filepath: str | Path, notes: list[str] | None = None
) -> dict[str, dict[str, ReferenceAnswer]]:
    """Parse one ``Эталон`` workbook.

    Returns ``{test_name: {normalized_question: ReferenceAnswer}}``. Skipped sheets
    are reported through ``notes`` (and the log) instead of vanishing silently.
    """
    path = Path(filepath)
    if not path.exists():
        raise ParseError("Файл не найден", filepath=str(path))

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(unreadable_file_message(path.name, exc), filepath=str(path)) from exc

    def note(msg: str) -> None:
        _log.warning(msg)
        if notes is not None:
            notes.append(msg)

    book: dict[str, dict[str, ReferenceAnswer]] = {}
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                try:
                    ws.reset_dimensions()
                    rows = list(ws.iter_rows(values_only=True))
                except Exception:  # noqa: BLE001
                    pass
            if all(_row_blank(r) for r in rows):
                continue  # genuinely empty
            if is_service_sheet(sheet_name):
                continue
            tlay = find_transposed_layout(rows)
            if tlay is not None:
                tmap = _parse_transposed(rows, tlay, sheet_name)
                if tmap:
                    book[sheet_name] = tmap
                else:
                    note(f"Лист эталона «{sheet_name}» пропущен: не найдено ни одного вопроса.")
                continue
            layout = find_reference_layout(rows)
            if layout is None:
                note(f"Лист эталона «{sheet_name}» пропущен: не найдены заголовки вопросов.")
                continue
            h, answer_start = layout
            header = rows[h]
            data_rows = rows[answer_start:]

            sheet_map: dict[str, ReferenceAnswer] = {}
            for col, q_cell in enumerate(header):
                question_text = norm_cell(q_cell)
                if not question_text:
                    continue
                # Collect every non-empty answer cell in this column.
                raw_values = [
                    norm_cell(r[col]) for r in data_rows if col < len(r) and norm_cell(r[col])
                ]
                raw_answer = " | ".join(raw_values)
                acceptable: list[str] = []
                for value in raw_values:
                    acceptable.extend(split_variants(value))
                sheet_map[normalize_question(question_text)] = ReferenceAnswer(
                    test_name=sheet_name,
                    question_text=question_text,
                    raw_answer=raw_answer,
                    acceptable=acceptable,
                )
            if sheet_map:
                book[sheet_name] = sheet_map
            else:
                note(f"Лист эталона «{sheet_name}» пропущен: не найдено ни одного вопроса.")
    finally:
        wb.close()

    return book
