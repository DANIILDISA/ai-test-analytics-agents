"""Which role does an input file play (answer array / reference key) and which course
does it belong to?  Name first, content second; folder first, file name second.

Kept separate from :mod:`pipeline.data_loader` so the heuristics are unit-testable.
"""

from __future__ import annotations

import re
from pathlib import Path

ETALON = "etalon"
MASSIV = "massiv"

_ETALON_WORDS = ("эталон", "etalon", "reference", "answer key", "answer_key")
_MASSIV_WORDS = ("массив", "array")

#: Folder names that are containers, not courses.
GENERIC_DIRS = {
    "data", "samples", "sample", "input", "inputs", "etalon", "reference", "uploads",
    "files", "tmp", "temp", "downloads", "desktop", "документы", "загрузки", "ответы",
}

_DIGITS_RE = re.compile(r"(?<![\d.])(\d{2,4})(?!\d)")


def fix_mojibake(name: str) -> str:
    """Repair a file name whose UTF-8 bytes were shown as cp1251 («РњР°СЃСЃРёРІ» ->
    «Массив»). Returns the name unchanged when it is not such mojibake."""
    try:
        fixed = name.encode("cp1251").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return name
    return fixed if fixed != name and re.search("[а-яА-ЯёЁ]", fixed) else name


def classify_by_name(name: str) -> str | None:
    """``etalon`` / ``massiv`` when the file name says so, else ``None``."""
    low = fix_mojibake(name).casefold()
    if any(w in low for w in _ETALON_WORDS):
        return ETALON
    if any(w in low for w in _MASSIV_WORDS):
        return MASSIV
    return None


def sniff_role(f: Path) -> str:
    """Decide the role from the file CONTENT (used when the name is silent).

    * xlsx: a sheet with the Тип/Результат/Полученный/Правильный sub-header row is an
      answer array; sheets with only a «Правильный ответ» label row are a reference.
    * csv / json: the field names.
    Defaults to ``massiv`` (the historical behaviour) when nothing can be told."""
    ext = f.suffix.lower()
    try:
        if ext == ".xlsx":
            import openpyxl

            from pipeline.parsers.answer_array_parser import _subheader_score
            from pipeline.parsers.reference_parser import _row_has_label, find_transposed_layout

            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            try:
                saw_label = False
                for ws in wb.worksheets:
                    for i, row in enumerate(ws.iter_rows(values_only=True)):
                        if i >= 25:
                            break
                        if _subheader_score(row) >= 3:
                            return MASSIV
                        if i >= 1 and _row_has_label(row):
                            saw_label = True
                    if find_transposed_layout(
                        [r for _, r in zip(range(25), ws.iter_rows(values_only=True))]
                    ):
                        saw_label = True
                return ETALON if saw_label else MASSIV
            finally:
                wb.close()
        if ext == ".csv":
            from pipeline.parsers.csv_parser import read_csv_text

            text, delim = read_csv_text(f)
            import csv as _csv
            import io as _io

            from pipeline.parsers.answer_array_parser import _subheader_score

            for i, r in enumerate(_csv.reader(_io.StringIO(text), delimiter=delim)):
                if i >= 25:
                    break
                if _subheader_score(tuple(r)) >= 3:
                    return MASSIV  # wide CSV: the same Тип/Результат/... sub-header row
            head = (text.splitlines() or [""])[0].casefold()
            cols = {c.strip() for c in head.split(delim)}
            answer_cols = {"result", "result_code", "given_answer", "given", "student_id",
                           "student", "пользователь", "результат", "полученный ответ"}
            return MASSIV if cols & answer_cols else ETALON
        if ext == ".json":
            import json

            data = json.loads(f.read_bytes().decode("utf-8-sig"))
            first = data[0] if isinstance(data, list) and data else {}
            if isinstance(first, dict) and ("answers" in first or "student_id" in first):
                return MASSIV
            return ETALON if isinstance(first, dict) and first else MASSIV
    except Exception:  # noqa: BLE001 - unreadable here; the real parser reports it properly
        pass
    return MASSIV


def classify(f: Path) -> str:
    """Role of an input file: by name first, by content when the name is silent."""
    by_name = classify_by_name(f.name)
    if by_name is not None:
        return by_name
    return sniff_role(f) if f.exists() else MASSIV


def course_id_for(f: Path, root: Path) -> str:
    """Course id = the parent folder name, unless that folder is only a container
    (``data``, ``samples``, ``uploads`` ...): then the first 2-4 digit group in the file
    name, falling back to the folder name."""
    pname = f.parent.name
    if re.search(r"\d{2,}", pname) or pname.casefold() not in GENERIC_DIRS:
        return pname
    m = _DIGITS_RE.search(f.stem)
    return m.group(1) if m else pname
