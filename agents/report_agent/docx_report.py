"""Word (.docx) report: clean A4 layout built with python-docx from ``report_model``.

Structure: title block → KPI tiles → «ИИ-выводы» (four blocks) → table of critical questions →
data remarks (only if any) → short methodology note. Typography: Calibri 10.5 pt, brand-blue
headings, generous spacing, bullet lists for list-like AI text, a repeating table header and
rows that never split across pages.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from agents.report_agent.report_model import build_report_model

BLUE = RGBColor(0x18, 0x53, 0xDC)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
RED = RGBColor(0xC2, 0x41, 0x0C)
FONT = "Calibri"


# ------------------------------------------------------------------ low-level helpers
def _font(obj, *, size=None, bold=None, color=None, italic=None) -> None:
    """Style a run or a style: font name (incl. Cyrillic/complex slots), size, weight, colour."""
    f = obj.font
    f.name = FONT
    el = getattr(obj, "_r", None)
    if el is None:
        el = obj.element
    rpr = el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def _shade(cell, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _cell_margins(cell, *, top=70, bottom=70, left=100, right=100) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def _table_borders(table, color="D1D5DB", size=4, inside=True) -> None:
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    sides = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for side in sides:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def _row_flags(row, *, header=False) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))
    if header:
        trpr.append(OxmlElement("w:tblHeader"))


def _bottom_border(paragraph, color="1853DC", size=8) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "2")
    b.set(qn("w:color"), color)
    pb.append(b)
    ppr.append(pb)


def _fmt(paragraph, *, before=0, after=6, line=1.15, keep_next=False) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next


def _para(doc_or_cell, text="", *, size=10.5, bold=False, color=INK, italic=False, before=0, after=6,
          align=None, keep_next=False, style=None):
    p = doc_or_cell.add_paragraph(style=style) if style else doc_or_cell.add_paragraph()
    if text:
        r = p.add_run(text)
        _font(r, size=size, bold=bold, color=color, italic=italic)
    _fmt(p, before=before, after=after, keep_next=keep_next)
    if align is not None:
        p.alignment = align
    return p


def _heading(doc, text: str, *, level=1):
    if level == 1:
        p = _para(doc, text, size=15, bold=True, color=BLUE, before=18, after=8, keep_next=True)
        _bottom_border(p)
    else:
        p = _para(doc, text, size=12, bold=True, color=INK, before=12, after=4, keep_next=True)
    return p


def _add_field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    _font(run, size=9, color=MUTED)
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


# ------------------------------------------------------------------ sections
def _title_block(doc, m: dict) -> None:
    _para(doc, m["title"].upper(), size=9.5, bold=True, color=BLUE, after=2)
    _para(doc, m["name"], size=22, bold=True, color=INK, after=4, keep_next=True)
    meta = [("Дата анализа", m["date"])]
    meta.append(("Группы", ", ".join(m["groups"]) if m["groups"] else "без разбивки на группы"))
    if m["scope"]:
        meta.append(("Охват", m["scope"]))
    for label, value in meta:
        p = doc.add_paragraph()
        _fmt(p, after=1)
        r1 = p.add_run(f"{label}: ")
        _font(r1, size=10.5, bold=True, color=MUTED)
        r2 = p.add_run(value)
        _font(r2, size=10.5, color=INK)


def _kpi_tiles(doc, m: dict) -> None:
    _heading(doc, "Ключевые показатели")
    table = doc.add_table(rows=1, cols=len(m["kpis"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _table_borders(table, color="E5E7EB", size=6, inside=True)
    width = Cm(16.6 / len(m["kpis"]))
    for cell, k in zip(table.rows[0].cells, m["kpis"]):
        cell.width = width
        _shade(cell, "F3F6FD")
        _cell_margins(cell, top=140, bottom=140, left=140, right=100)
        cell.paragraphs[0].text = ""
        v = cell.paragraphs[0]
        _fmt(v, after=2, line=1.0)
        _font(v.add_run(k["value"]), size=22, bold=True, color=BLUE)
        lab = cell.add_paragraph()
        _fmt(lab, after=0, line=1.05)
        _font(lab.add_run(k["label"]), size=9, color=MUTED)
        if k.get("hint"):
            h = cell.add_paragraph()
            _fmt(h, after=0, line=1.05)
            _font(h.add_run(k["hint"]), size=8.5, color=MUTED)
    _row_flags(table.rows[0])


def _ai_section(doc, m: dict) -> None:
    _heading(doc, "ИИ-выводы")
    if not m["ai_blocks"]:
        _para(doc, "Глубокий анализ с ИИ для этого анализа не запускался, поэтому ИИ-выводов нет. "
                   "Числовые результаты — в таблице ниже и в Excel-выгрузке.", color=MUTED)
        return
    for block in m["ai_blocks"]:
        _heading(doc, block["title"], level=2)
        for kind, text in block["items"]:
            if kind == "li":
                p = doc.add_paragraph(style="List Bullet")
                _fmt(p, after=4)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.4)
                _font(p.add_run(text), size=10.5, color=INK)
            else:
                _para(doc, text, after=6)


def _critical_table(doc, m: dict) -> None:
    _heading(doc, "Критические вопросы")
    if not m["critical"]:
        _para(doc, "Критических вопросов нет: ни по одному вопросу доля верных ответов не опустилась ниже порога.",
              color=MUTED)
        return
    _para(doc, f"Вопросов с низкой долей верных ответов: {m['critical_total']}. "
               "Сначала самые трудные.", color=MUTED, after=6)
    heads = ["Вопрос", "Верно", "Типичный неверный ответ", "Верный ответ"]
    widths = [Cm(7.0), Cm(1.5), Cm(4.3), Cm(3.8)]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _table_borders(table)
    for cell, text, w in zip(table.rows[0].cells, heads, widths):
        cell.width = w
        _shade(cell, "1853DC")
        _cell_margins(cell, top=80, bottom=80)
        cell.paragraphs[0].text = ""
        p = cell.paragraphs[0]
        _fmt(p, after=0, line=1.0)
        _font(p.add_run(text), size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    _row_flags(table.rows[0], header=True)

    for i, q in enumerate(m["critical"]):
        row = table.add_row()
        _row_flags(row)
        values = [
            q["question"],
            q["rate"],
            q["wrong"] or "—",
            (q["correct"] or "нет в эталоне"),
        ]
        for j, (cell, text, w) in enumerate(zip(row.cells, values, widths)):
            cell.width = w
            if i % 2:
                _shade(cell, "F9FAFB")
            _cell_margins(cell)
            cell.paragraphs[0].text = ""
            p = cell.paragraphs[0]
            _fmt(p, after=0, line=1.1)
            bold = j == 1
            color = RED if j == 1 else (MUTED if (j == 3 and not q["correct"]) else INK)
            _font(p.add_run(text), size=9.5, bold=bold, color=color)
            if j == 3 and q["key_suspect"]:
                n = cell.add_paragraph()
                _fmt(n, before=2, after=0, line=1.05)
                _font(n.add_run("Возможна ошибка в эталоне: проверьте ключ."), size=8.5, italic=True, color=RED)
    if m["critical_more"]:
        _para(doc, f"Ещё вопросов: {m['critical_more']} — полный список в Excel-выгрузке.",
              size=9.5, color=MUTED, before=6, after=0)


def _notes_and_method(doc, m: dict) -> None:
    if m["data_notes"]:
        _heading(doc, "Замечания к данным")
        for t in m["data_notes"]:
            p = doc.add_paragraph(style="List Bullet")
            _fmt(p, after=3)
            _font(p.add_run(t), size=10, color=INK)
    _heading(doc, "Методика")
    for t in m["method"]:
        p = doc.add_paragraph(style="List Bullet")
        _fmt(p, after=2)
        _font(p.add_run(t), size=9.5, color=MUTED)


# ------------------------------------------------------------------ entry point
def build_docx(result: dict, meta: dict | None = None) -> bytes:
    """Render a saved run to a .docx (bytes)."""
    m = build_report_model(result, meta)
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    _font(normal, size=10.5, color=INK)
    for style_name in ("List Bullet",):
        _font(doc.styles[style_name], size=10.5, color=INK)

    props = doc.core_properties
    props.title = f"{m['title']} — {m['name']}"
    props.author = "КОРПУС C2"
    props.language = "ru-RU"

    # footer: name + page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(fp.add_run(f"{m['name']} · стр. "), size=9, color=MUTED)
    _add_field(fp, "PAGE")

    _title_block(doc, m)
    _kpi_tiles(doc, m)
    _ai_section(doc, m)
    _critical_table(doc, m)
    _notes_and_method(doc, m)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
